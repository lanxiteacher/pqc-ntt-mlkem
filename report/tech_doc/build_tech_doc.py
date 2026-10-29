"""Fill the official technical-document template from technical_document.md.

Usage:
    python report/tech_doc/build_tech_doc.py              # build DOCX
    python report/tech_doc/build_tech_doc.py --check-only # word limits only

Output: _submission/docs/technical_document.docx (gitignored).
The template 技术文档（模板）.docx must sit at the repository root.
"""
import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
MD = Path(__file__).with_name('technical_document.md')
TEMPLATE_CANDIDATES = [ROOT / '技术文档（模板）.docx', Path(__file__).with_name('template.docx')]
OUT = ROOT / '_submission' / 'docs' / 'technical_document.docx'

EDGE_CANDIDATES = [
    Path(r"C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe"),
    Path(r"C:/Program Files/Microsoft/Edge/Application/msedge.exe"),
    Path(r"C:/Program Files/Google/Chrome/Application/chrome.exe"),
]

LIMITS = {
    '摘要': 800, '功能与特性': 400, '应用领域': 400, '主要技术特点': 400,
    '主要性能指标': 200, '主要创新点': 200, '设计流程': 200,
    '可扩展之处': 300, '心得体会': 1000,
}

ANCHORS = [
    (0, '作品名称', 'TITLE'),
    (3, '800字内', '摘要'),
    (7, '400字内（可加重要图）', '功能与特性'),
    (10, '400字内（可加重要图）', '应用领域'),
    (13, '400字内；', '主要技术特点'),
    (16, '200字内（建议用表格）', '主要性能指标'),
    (19, '200字以内（逐点给出）', '主要创新点'),
    (22, '200字内（可加重要图）', '设计流程'),
    (28, '给出系统整体框图', '系统整体介绍'),
    (31, '2.2.1 硬件整体介绍', '2.2.1 硬件整体介绍'),
    (32, '2.2.2 机械设计介绍', '2.2.2 机械设计介绍'),
    (33, '2.2.3 电路各模块介绍', '2.2.3 电路各模块介绍'),
    (36, '2.3.1 软件整体介绍', '2.3.1 软件整体介绍'),
    (37, '2.3.2 软件各模块介绍', '2.3.2 软件各模块介绍'),
    (41, '整体介绍（整个系统实物', '3.1 整体介绍'),
    (44, '3.2.1 机械成果', '3.2.1 机械成果'),
    (45, '3.2.2 电路成果', '3.2.2 电路成果'),
    (46, '3.2.3 软件成果', '3.2.3 软件成果'),
    (48, '特性成果', '3.3 特性成果'),
    (54, '300字内', '可扩展之处'),
    (56, '1000字内', '心得体会'),
    (59, '按照标准格式', '参考文献'),
]


def char_count(text):
    cjk = len(re.findall(r'[一-鿿]', text))
    words = len(re.findall(r'[A-Za-z0-9][A-Za-z0-9_.%µ/+\-]*', text))
    return cjk, words, cjk + words


def parse_md(path):
    text = path.read_text(encoding='utf-8')
    title = re.search(r'^# (.+)$', text, re.M).group(1).strip()
    chunks = re.split(r'^## (.+)$', text, flags=re.M)[1:]
    sections = {}
    for name, body in zip(chunks[::2], chunks[1::2]):
        sections[name.strip()] = body
    return title, sections


def block_text(body):
    """Plain text of a section, used for word counting."""
    lines = []
    for line in body.splitlines():
        s = line.strip()
        if not s or s.startswith('|') or s.startswith('!['):
            continue
        lines.append(s)
    return '\n'.join(lines)


def parse_blocks(body):
    blocks, lines, i = [], body.splitlines(), 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        m = re.match(r'!\[(.*?)\]\((.*?)\)\s*$', line.strip())
        if m:
            blocks.append(('img', m.group(2), m.group(1)))
            i += 1
            continue
        if line.lstrip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(('|', '![')):
            para.append(lines[i].strip())
            i += 1
        blocks.append(('p', ''.join(para)))
    return blocks


def find_edge():
    for p in EDGE_CANDIDATES:
        if p.exists():
            return p
    return None


def svg_to_png(svg, tmpdir, edge):
    txt = svg.read_text(encoding='utf-8', errors='replace')
    m = re.search(r'viewBox="([\d.\s-]+)"', txt)
    if not m:
        raise RuntimeError(f'no viewBox in {svg}')
    _, _, w, h = [float(x) for x in m.group(1).split()]
    W, H = int(w * 2), int(h * 2)
    html = tmpdir / (svg.stem + '.html')
    html.write_text(
        f'<html><body style="margin:0"><img src="{svg.as_uri()}" '
        f'style="width:{W}px;height:{H}px"></body></html>', encoding='utf-8')
    out = tmpdir / (svg.stem + '.png')
    subprocess.run([str(edge), '--headless=new', '--disable-gpu', '--hide-scrollbars',
                    f'--window-size={W},{H}', f'--screenshot={out}', html.as_uri()],
                   check=True, timeout=180)
    return out


def prepare_image(path, tmpdir, edge):
    path = (ROOT / path).resolve()
    if path.suffix.lower() == '.svg':
        return svg_to_png(path, tmpdir, edge)
    im = Image.open(path)
    if im.width <= 1600:
        return path
    scale = 1600 / im.width
    out = tmpdir / (path.stem + '_w1600.png')
    im.resize((1600, int(im.height * scale)), Image.LANCZOS).save(out)
    return out


def fmt_run(run, size=12):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), '宋体')


def set_body_format(p, first_indent=True):
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Pt(0)
    pf.first_line_indent = Pt(24) if first_indent else Pt(0)


def clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    pPr = p._p.pPr
    if pPr is not None:
        for tag in ('w:ind', 'w:jc'):
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)


def set_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def add_image(doc, cursor, path, caption, tmpdir, edge):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    with Image.open(path) as im:
        portrait = im.height > im.width * 1.05
    run = p.add_run()
    if 'gui_act' in path.name:
        run.add_picture(str(path), width=Cm(12.5))
    elif portrait:
        run.add_picture(str(path), height=Cm(11.5))
    else:
        run.add_picture(str(path), width=Cm(13.5))
    cursor.addnext(p._p)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    fmt_run(r, 10.5)
    p._p.addnext(cap._p)
    return cap._p


def add_table(doc, cursor, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = True
    set_borders(table)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    header = table.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    header.append(el)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            fmt_run(r, 10.5)
            r.bold = i == 0
    cursor.addnext(table._tbl)
    return table._tbl


def fill_anchor(doc, anchor, blocks, tmpdir, edge):
    clear_runs(anchor)
    set_body_format(anchor)
    cursor = anchor._p
    anchor_free = False
    for block in blocks:
        if block[0] == 'img':
            cursor = add_image(doc, cursor, prepare_image(block[1], tmpdir, edge),
                               block[2], tmpdir, edge)
        elif block[0] == 'table':
            cursor = add_table(doc, cursor, block[1])
        else:
            if anchor_free:
                p = doc.add_paragraph()
                set_body_format(p)
                cursor.addnext(p._p)
            else:
                p = anchor
                anchor_free = True
            r = p.add_run(block[1])
            fmt_run(r)
            cursor = p._p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check-only', action='store_true')
    args = ap.parse_args()

    title, sections = parse_md(MD)
    bad = []
    for name, limit in LIMITS.items():
        cjk, words, total = char_count(block_text(sections[name]))
        flag = 'OK' if total <= limit else f'OVER {total - limit}'
        print(f'{name:8s} {cjk:4d} 汉字 + {words:3d} 词 = {total:4d} / {limit}  {flag}')
        if total > limit:
            bad.append(name)
    refs = len(re.findall(r'^\[\d+\]', sections['参考文献'], re.M))
    print(f'参考文献 {refs} 条 / 20')
    if refs > 20:
        bad.append('参考文献')
    if bad:
        sys.exit('超出上限: ' + ', '.join(bad))
    if args.check_only:
        return

    template = next((t for t in TEMPLATE_CANDIDATES if t.exists()), None)
    if template is None:
        sys.exit('找不到模板：根目录或 report/tech_doc/template.docx')
    edge = find_edge()
    if edge is None:
        sys.exit('找不到 Edge/Chrome，无法转换 SVG 图')

    doc = Document(str(template))
    paragraphs = list(doc.paragraphs)
    # Title block: replace name, drop the "no school name" note.
    p0 = paragraphs[0]
    clear_runs(p0)
    r = p0.add_run(title)
    fmt_run(r, 16)
    r.bold = True
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[1]._p.getparent().remove(paragraphs[1]._p)

    # Keep long tables whole: start these sections on a fresh page.
    for idx in (37, 39, 48):
        paragraphs[idx].paragraph_format.page_break_before = True

    with tempfile.TemporaryDirectory(prefix='tech_doc_') as td:
        tmpdir = Path(td)
        for idx, expect, key in ANCHORS:
            if key == 'TITLE':
                continue
            anchor = paragraphs[idx]
            if expect not in anchor.text:
                sys.exit(f'模板第 {idx} 段不是「{expect}」：{anchor.text[:40]!r}')
            fill_anchor(doc, anchor, parse_blocks(sections[key]), tmpdir, edge)
        OUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(OUT))

    check = Document(str(OUT))
    print(f'已生成 {OUT.relative_to(ROOT)}')
    print(f'段落 {len(check.paragraphs)}；表格 {len(check.tables)}；图片 {len(check.inline_shapes)}')


if __name__ == '__main__':
    main()
